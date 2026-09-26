from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Mapping, Sequence, Union

from openpyxl import load_workbook


def _normalize(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_label(value: str | None) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split()).lower()


def _parse_node_path(value: Union[str, Sequence[str], None]) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, str):
        parts = [part.strip() for part in value.split(">")]
        return tuple(part for part in parts if part)
    return tuple(str(part).strip() for part in value if str(part).strip())


def _canonical_path_key(value: Union[str, Sequence[str], None]) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, str):
        parts = [part.strip() for part in value.split(">") if part.strip()]
    else:
        parts = [str(part).strip() for part in value if str(part).strip()]
    return tuple(_normalize_label(part) for part in parts)


def build_incident_graph(source: Union[str, Path]):
    """
    Build a tree-like graph from Excel columns F:J.

    The workbook typically stores values in the following order:
    F = group of incident / top level
    G = first subdivision
    H = second subdivision
    I = third subdivision
    J = last available attribute

    Example path:
        ["Пожары и задымления", "на улице", "мусор", "открытое пламя"]
    """
    path = Path(source)
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook[workbook.sheetnames[0]]

    adjacency: dict[tuple[str, ...], list[str]] = defaultdict(list)
    node_counts: dict[tuple[str, ...], int] = defaultdict(int)

    for row in ws.iter_rows(min_row=1, values_only=True):
        values = [_normalize(value) for value in row[5:10]]
        values = [value for value in values if value]
        if not values:
            continue

        current_path: list[str] = []
        for value in values:
            if not current_path or value != current_path[-1]:
                current_path.append(value)

        if not current_path:
            continue

        for depth in range(len(current_path)):
            parent = tuple(current_path[:depth])
            child = current_path[depth]
            if child not in adjacency[parent]:
                adjacency[parent].append(child)
            node_counts[(parent + (child,))] += 1

    root_children = adjacency.get((), [])
    normalized = {key: list(value) for key, value in adjacency.items()}
    return {
        "root": (),
        "children": normalized,
        "root_children": root_children,
        "node_counts": {
            k: v for k, v in sorted(node_counts.items(), key=lambda item: len(item[0]))
        },
    }


def _path_to_key(path: Sequence[str]) -> str:
    return ">".join(path)


def _key_to_path(key: str) -> tuple[str, ...]:
    return tuple(part.strip() for part in key.split(">") if part.strip())


def get_next_levels(graph: Mapping, path: Union[str, Sequence[str], None]) -> list[str]:
    """
    Return all available next-level values for the given path.

    Examples:
        get_next_levels(graph, []) -> root-level incident groups
        get_next_levels(graph, ["Пожары и задымления"]) -> next categories
        get_next_levels(graph, "Пожары и задымления > на улице") -> further options
    """
    current = _parse_node_path(path)
    children = graph.get("children", {})

    current_canonical = _canonical_path_key(current)
    for key, value in children.items():
        key_canonical = _canonical_path_key(key)
        if key_canonical == current_canonical:
            return list(value)

    if not current:
        return list(graph.get("root_children", []))

    return []


def _serialize_graph(graph: Mapping) -> dict:
    serializable = {
        "root": list(graph.get("root", ())),
        "root_children": list(graph.get("root_children", [])),
        "children": {},
        "node_counts": {},
    }

    for key, value in graph.get("children", {}).items():
        serializable["children"][_path_to_key(key)] = list(value)

    for key, value in graph.get("node_counts", {}).items():
        serializable["node_counts"][_path_to_key(key)] = value

    return serializable


def _deserialize_graph(graph: Mapping) -> dict:
    restored = {
        "root": tuple(graph.get("root", ())),
        "root_children": list(graph.get("root_children", [])),
        "children": {},
        "node_counts": {},
    }

    for key, value in graph.get("children", {}).items():
        restored["children"][_key_to_path(str(key))] = list(value)

    for key, value in graph.get("node_counts", {}).items():
        restored["node_counts"][_key_to_path(str(key))] = value

    return restored


def save_incident_graph(graph: Mapping, output_path: Union[str, Path]):
    """
    Save the prebuilt graph to JSON. This is the recommended production pattern:
    build the graph once, cache it, and then load only the graph file when needed.
    """
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    serializable = _serialize_graph(graph)
    output.write_text(
        json.dumps(serializable, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return str(output)


def load_incident_graph(path: Union[str, Path]) -> dict:
    """Load a saved graph JSON file."""
    file_path = Path(path)
    with file_path.open("r", encoding="utf-8") as handle:
        loaded = json.load(handle)
    return _deserialize_graph(loaded)


def export_graph_json(source: Union[str, Path], output_path: Union[str, Path]):
    graph = build_incident_graph(source)
    return save_incident_graph(graph, output_path)


def _build_argument_parser():
    parser = argparse.ArgumentParser(
        description="Load Excel incident classification into a graph-like structure."
    )
    parser.add_argument(
        "excel_path", type=str, help="Path to the Excel file with columns F:J"
    )
    parser.add_argument(
        "-o", "--output", type=str, default=None, help="Optional output JSON path"
    )
    return parser


if __name__ == "__main__":
    args = _build_argument_parser().parse_args()
    graph = build_incident_graph(args.excel_path)
    print(
        json.dumps(
            {
                "root_children": graph["root_children"][:10],
                "total_root_groups": len(graph["root_children"]),
                "example_next_levels": (
                    get_next_levels(graph, graph["root_children"][0])
                    if graph["root_children"]
                    else []
                ),
            },
            ensure_ascii=False,
            indent=2,
        )
    )

    if args.output:
        export_graph_json(args.excel_path, args.output)
        print(f"Graph exported to {args.output}")
